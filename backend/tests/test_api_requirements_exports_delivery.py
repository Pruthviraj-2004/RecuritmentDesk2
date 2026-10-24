import io
import uuid
from datetime import timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.types import utcnow
from app.integrations.gmail_client import GmailAmbiguousSend, GmailAuthError, GmailTransient
from app.models.email import EmailAccount
from app.models.ops import AuditLog, EmailDelivery, ExportJob, ProcessingJob
from app.models.requirements import JobRequirement
from tests.conftest import login, make_user

API = "/api/v1"


def mk(api, **kw):
    body = {"job_title": "Python Developer", "client_name": "Globex", "required_skills": ["Python", "Django"], "location": "Pune", "work_mode": "remote",
            "minimum_experience": 3, "maximum_experience": 6, "recruiter_email": "jane@staffing.com", **kw}
    r = api.post(f"{API}/requirements", json=body)
    assert r.status_code == 201, r.text
    return r.json()


# ---------------------------------------------------------------- requirements CRUD, filters, versions
def test_create_edit_versions_and_validation(api):
    r = mk(api)
    rid = r["id"]
    assert r["review_status"] == "needs_review" and r["required_skills"] == ["Django", "Python"] and r["version_number"] == 1
    p = api.patch(f"{API}/requirements/{rid}", json={"job_title": "Lead Python Developer", "salary_min": 100, "salary_max": 200, "preferred_skills": ["Redis"]})
    assert p.status_code == 200 and p.json()["version_number"] == 2 and p.json()["preferred_skills"] == ["Redis"]
    assert api.patch(f"{API}/requirements/{rid}", json={"salary_min": 300, "salary_max": 200}).status_code == 422
    assert api.patch(f"{API}/requirements/{rid}", json={"recruiter_email": "not-an-email"}).status_code == 422
    assert api.patch(f"{API}/requirements/{rid}", json={"id": "x"}).status_code in (200, 422)
    v = api.get(f"{API}/requirements/{rid}/versions").json()
    assert [x["version_number"] for x in v] == [2, 1] and v[0]["previous_values"]["job_title"] == "Python Developer"
    # an unchanged PATCH does not create a version
    assert api.patch(f"{API}/requirements/{rid}", json={"job_title": "Lead Python Developer"}).json()["version_number"] == 2


def test_status_transitions_bulk_partial_failure_and_soft_delete(api):
    a, b = mk(api, job_title="A")["id"], mk(api, job_title="B")["id"]
    assert api.post(f"{API}/requirements/{a}/approve").json()["review_status"] == "approved"
    assert api.post(f"{API}/requirements/{a}/reject").json()["review_status"] == "rejected"
    assert api.post(f"{API}/requirements/{a}/archive").json()["lifecycle_status"] == "archived"
    r = api.post(f"{API}/requirements/{a}/restore").json()
    assert r["lifecycle_status"] == "active" and r["review_status"] == "needs_review"
    res = api.post(f"{API}/requirements/bulk", json={"ids": [a, b, "00000000-0000-0000-0000-000000000001"], "action": "approve"}).json()
    assert sorted(res["succeeded"]) == sorted([a, b]) and len(res["failed"]) == 1
    assert api.delete(f"{API}/requirements/{b}").json() == {"archived": True}  # soft delete
    assert api.get(f"{API}/requirements/{b}").json()["lifecycle_status"] == "archived"


def test_list_search_filters_sort_and_pagination(api):
    mk(api, job_title="Python Developer", client_name="Globex", location="Pune", minimum_experience=2, maximum_experience=4, salary_min=10, salary_max=20)
    mk(api, job_title="Java Architect", client_name="Initech", location="Delhi", work_mode="onsite", required_skills=["Java", "Spring Boot"], minimum_experience=9,
       maximum_experience=12, recruiter_email="bob@x.com")
    for i in range(3):
        mk(api, job_title=f"Data Analyst {i}", client_name="Hooli", location="Mumbai", required_skills=["SQL"])
    g = lambda **p: api.get(f"{API}/requirements", params=p).json()  # noqa: E731
    assert g()["total"] == 5
    assert g(q="java")["total"] == 1 and g(q="bob@x.com")["total"] == 1 and g(q="spring")["total"] == 1
    assert g(client="hool")["total"] == 3 and g(skill="sql")["total"] == 3 and g(location="del")["total"] == 1
    assert g(work_mode="onsite")["total"] == 1 and g(exp_min=10)["total"] == 1 and g(salary_max=15)["total"] == 1
    assert g(q="100%")["total"] == 0  # LIKE wildcards are escaped
    p1 = g(page_size=2, page=1, sort="job_title", direction="asc")
    assert [i["job_title"] for i in p1["items"]] == ["Data Analyst 0", "Data Analyst 1"] and p1["total"] == 5
    assert g(page_size=2, page=3, sort="job_title", direction="asc")["items"][0]["job_title"] == "Python Developer"
    assert g(review_status="approved")["total"] == 0 and g(duplicate_status="pending")["total"] == 0
    assert api.get(f"{API}/requirements", params={"page_size": 500}).status_code == 422


def test_similar_requirements(api):
    a = mk(api, job_title="Python Developer")["id"]
    mk(api, job_title="Python Developer", location="Pune")
    mk(api, job_title="Cook", client_name="Other", required_skills=["Cooking"])
    sims = api.get(f"{API}/requirements/{a}/similar").json()
    assert sims and sims[0]["requirement"]["job_title"] == "Python Developer" and sims[0]["score"] > 0.5
    assert all(s["requirement"]["id"] != a for s in sims)


def test_requirements_are_isolated_between_organizations(client, api, db):
    rid = mk(api)["id"]
    make_user(db, "other@example.com", org="Org2")
    o = login(client, "other@example.com")
    assert o.get(f"{API}/requirements").json()["total"] == 0
    assert o.get(f"{API}/requirements/{rid}").status_code == 404
    assert o.patch(f"{API}/requirements/{rid}", json={"job_title": "hack"}).status_code == 404
    assert o.post(f"{API}/requirements/{rid}/approve").status_code == 404
    assert o.post(f"{API}/exports", json={"scope": "selected", "selected_ids": [rid]}).json()["status"] == "failed"  # nothing from other orgs is exported


# ---------------------------------------------------------------- exports
def make_export(api, **kw):
    r = api.post(f"{API}/exports", json=kw)
    assert r.status_code == 202, r.text
    return r.json()


def test_export_selected_and_download_workbook(api, db):
    ids = [mk(api, job_title=f"Role {i}")["id"] for i in range(3)]
    api.post(f"{API}/requirements/{ids[0]}/approve")
    e = make_export(api, scope="selected", selected_ids=[ids[0], ids[1], ids[0]], filename="My Report!", options={"columns": ["job_title", "client_name"], "sheet_name": "Hiring"})
    assert e["status"] == "completed" and e["record_count"] == 2 and e["filename"] == "My-Report.xlsx"
    d = api.get(f"{API}/exports/{e['id']}/download")
    assert d.status_code == 200 and "spreadsheetml" in d.headers["content-type"] and d.headers["cache-control"] == "private, no-store"
    wb = load_workbook(io.BytesIO(d.content))
    ws = wb["Hiring"]
    assert [c.value for c in ws[1]] == ["Job Title", "Client"] and ws.max_row == 3
    assert "Summary" in wb.sheetnames


def test_export_scopes_filters_and_validation(api):
    a, b = mk(api, job_title="A", client_name="X")["id"], mk(api, job_title="B", client_name="Y")["id"]
    api.post(f"{API}/requirements/{a}/approve")
    assert make_export(api, scope="approved")["record_count"] == 1
    assert make_export(api, scope="filtered", filters={"client": "Y"})["record_count"] == 1
    assert make_export(api, scope="filtered")["record_count"] == 2
    assert api.post(f"{API}/exports", json={"scope": "selected"}).status_code == 422
    assert api.post(f"{API}/exports", json={"scope": "filtered", "options": {"columns": ["nope"]}}).status_code == 422
    empty = make_export(api, scope="filtered", filters={"client": "ZZZ"})
    assert empty["status"] == "failed" and "No requirements" in empty["error_summary"]
    assert api.get(f"{API}/exports").json()["total"] == 4


def test_export_excludes_merged_and_uses_source_columns_when_asked(api, db):
    r = mk(api)["id"]
    e = make_export(api, scope="filtered", options={"include_source": True, "include_audit": True, "include_metadata": True})
    wb = load_workbook(io.BytesIO(api.get(f"{API}/exports/{e['id']}/download").content))
    headers = [c.value for c in wb.worksheets[0][1]]
    assert "Source Emails" in headers and "Review Status" in headers and "Export Info" in wb.sheetnames
    row = db.scalar(select(JobRequirement).where(JobRequirement.id == uuid.UUID(r)))
    row.lifecycle_status = "merged"
    db.commit()
    assert make_export(api, scope="selected", selected_ids=[r])["status"] == "failed"


def test_export_download_requires_auth_org_and_expiry(client, api, db):
    mk(api)
    e = make_export(api, scope="filtered")
    anon = client.__class__(client.app)
    assert anon.get(f"{API}/exports/{e['id']}/download").status_code == 401
    make_user(db, "other@example.com", org="Org2")
    assert login(client.__class__(client.app), "other@example.com").get(f"{API}/exports/{e['id']}/download").status_code == 404
    exp = db.scalar(select(ExportJob))
    exp.expires_at = utcnow() - timedelta(hours=1)
    db.commit()
    r = api.get(f"{API}/exports/{e['id']}/download")
    assert r.status_code == 403 and r.json()["error"]["code"] == "export_expired"
    from app.services.exports.runner import cleanup_expired_exports

    assert cleanup_expired_exports(db) == 1 and db.scalar(select(ExportJob)).status == "expired"


def test_formula_payload_in_data_is_neutralised_in_export(api):
    mk(api, client_name='=HYPERLINK("http://evil","click")', job_title="+cmd|' /C calc'!A0")
    e = make_export(api, scope="filtered")
    ws = load_workbook(io.BytesIO(api.get(f"{API}/exports/{e['id']}/download").content)).worksheets[0]
    vals = [c.value for c in ws[2] if c.value]
    assert vals and all(not str(v).startswith(("=", "+", "-", "@")) for v in vals)
    assert all(c.data_type != "f" for c in ws[2])


# ---------------------------------------------------------------- delivery
def setup_export(api):
    mk(api)
    return make_export(api, scope="filtered")


def payload(acct, **kw):
    return {"sender_account_id": str(acct.id), "to": ["client@example.com"], "cc": ["boss@example.com"], "subject": "Weekly requirements", "body": "Hi", "confirm": True, **kw}


def test_send_requires_confirmation_valid_recipients_and_send_scope(api, db, connected_account, mailbox):
    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    url = f"{API}/exports/{e['id']}/send"
    r = api.post(url, json=payload(a, confirm=False))
    assert r.status_code == 422 and r.json()["error"]["code"] == "confirmation_required"
    assert api.post(url, json=payload(a, to=["bad address"])).status_code == 422
    assert api.post(url, json=payload(a, to=["a@b.com\r\nBcc: x@y.com"])).status_code == 422
    assert mailbox.for_email("sender@gmail.com").sent == [] and db.scalar(select(EmailDelivery)) is None  # nothing was sent
    ro = connected_account("readonly@gmail.com", scopes=["https://www.googleapis.com/auth/gmail.readonly"])
    r = api.post(url, json=payload(ro))
    assert r.status_code == 409 and r.json()["error"]["code"] == "missing_send_scope"


def test_send_success_attaches_workbook_and_is_idempotent(api, db, connected_account, mailbox):
    from email import message_from_bytes, policy

    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    r = api.post(f"{API}/exports/{e['id']}/send", json=payload(a), headers={"Idempotency-Key": "abc-123"})
    assert r.status_code == 202
    d = r.json()
    assert d["status"] == "accepted_by_gmail" and "not guaranteed" in d["status_label"] and d["provider_message_id"] and d["sender_email"] == "sender@gmail.com"
    box = mailbox.for_email("sender@gmail.com")
    assert len(box.sent) == 1
    msg = message_from_bytes(box.sent[0], policy=policy.default)
    assert msg["To"] == "client@example.com" and msg["Cc"] == "boss@example.com" and msg["Subject"] == "Weekly requirements"
    att = list(msg.iter_attachments())[0]
    assert att.get_filename().endswith(".xlsx") and load_workbook(io.BytesIO(att.get_payload(decode=True))).worksheets[0].max_row == 2
    again = api.post(f"{API}/exports/{e['id']}/send", json=payload(a), headers={"Idempotency-Key": "abc-123"})
    assert again.json()["id"] == d["id"] and len(box.sent) == 1  # no double send
    assert api.get(f"{API}/deliveries").json()["total"] == 1
    actions = [x.action for x in db.scalars(select(AuditLog))]
    assert "delivery.requested" in actions and "delivery.accepted" in actions


def test_failed_send_is_reported_and_retry_works(api, db, connected_account, mailbox):
    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    box = mailbox.for_email("sender@gmail.com")
    box.send_error = GmailTransient("503")
    d = api.post(f"{API}/exports/{e['id']}/send", json=payload(a)).json()
    assert d["status"] == "failed" and "not sent" in d["error_summary"].lower() and box.sent == []
    box.send_error = None
    r = api.post(f"{API}/deliveries/{d['id']}/retry", json={})
    assert r.status_code == 202 and r.json()["status"] == "accepted_by_gmail" and len(box.sent) == 1
    assert api.post(f"{API}/deliveries/{d['id']}/retry", json={}).status_code == 409  # accepted deliveries cannot be retried


def test_ambiguous_network_failure_never_auto_resends(api, db, connected_account, mailbox):
    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    box = mailbox.for_email("sender@gmail.com")
    box.send_error = GmailAmbiguousSend("timeout")
    d = api.post(f"{API}/exports/{e['id']}/send", json=payload(a)).json()
    assert d["status"] == "uncertain" and "Sent folder" in d["error_summary"] and box.sent == []
    box.send_error = None
    r = api.post(f"{API}/deliveries/{d['id']}/retry", json={})
    assert r.status_code == 409 and r.json()["error"]["code"] == "uncertain_delivery_ack_required"
    assert box.sent == []


def test_ambiguous_failure_but_message_actually_accepted_is_detected(api, db, connected_account, mailbox):
    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    box = mailbox.for_email("sender@gmail.com")
    box.send_error, box.send_accepted_despite_error = GmailAmbiguousSend("timeout"), True
    d = api.post(f"{API}/exports/{e['id']}/send", json=payload(a)).json()
    assert d["status"] == "accepted_by_gmail" and len(box.sent) == 1  # found in Sent via Message-ID, not re-sent


def test_retry_after_uncertain_checks_sent_folder_before_resending(api, db, connected_account, mailbox):
    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    box = mailbox.for_email("sender@gmail.com")
    box.send_error = GmailAmbiguousSend("timeout")
    d = api.post(f"{API}/exports/{e['id']}/send", json=payload(a)).json()
    # the message was in fact delivered after the client gave up
    row = db.scalar(select(EmailDelivery))
    box.sent.append(f"Message-ID: {row.rfc_message_id}".encode())
    box.send_error = None
    r = api.post(f"{API}/deliveries/{d['id']}/retry", json={"acknowledge_possible_duplicate": True})
    assert r.json()["status"] == "accepted_by_gmail" and len(box.sent) == 1


def test_auth_failure_expires_sender_account(api, db, connected_account, mailbox):
    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    mailbox.for_email("sender@gmail.com").send_error = GmailAuthError("401")
    d = api.post(f"{API}/exports/{e['id']}/send", json=payload(a)).json()
    assert d["status"] == "failed" and "reauthorize" in d["error_summary"].lower()
    assert db.get(EmailAccount, a.id).connection_status == "expired"
    r = api.post(f"{API}/exports/{e['id']}/send", json=payload(a))
    assert r.status_code == 409 and r.json()["error"]["code"] == "account_needs_reauth"


def test_attachment_size_limit_and_authorization(client, api, db, connected_account, mailbox, monkeypatch):
    from app.core.config import get_settings

    a = connected_account("sender@gmail.com")
    e = setup_export(api)
    monkeypatch.setattr(get_settings(), "max_delivery_attachment_mb", 0)
    r = api.post(f"{API}/exports/{e['id']}/send", json=payload(a))
    assert r.status_code == 422 and r.json()["error"]["code"] == "attachment_too_large"
    monkeypatch.undo()
    other_user, _ = make_user(db, "other@example.com", org="Org2")
    o = login(client, "other@example.com")
    assert o.post(f"{API}/exports/{e['id']}/send", json=payload(a)).status_code in (404, 409, 422)
    assert mailbox.for_email("sender@gmail.com").sent == []
    anon = client.__class__(client.app)
    assert anon.post(f"{API}/exports/{e['id']}/send", json=payload(a)).status_code == 401


# ---------------------------------------------------------------- jobs, audit, settings, analytics, health
def test_jobs_cancel_retry_and_summary(api, db, connected_account, mailbox):
    from app.services import jobs

    a = connected_account("one@gmail.com")
    org_id = a.organization_id
    j = jobs.create_job(db, org_id=org_id, job_type="sync", account_id=a.id)
    db.commit()
    assert api.post(f"{API}/jobs/{j.id}/cancel").json()["status"] == "cancelled"
    assert api.post(f"{API}/jobs/{j.id}/cancel").status_code == 409
    r = api.post(f"{API}/jobs/{j.id}/retry")
    assert r.status_code == 200 and r.json()["id"] != str(j.id)
    lst = api.get(f"{API}/jobs").json()
    assert lst["total"] == 2 and api.get(f"{API}/jobs", params={"status": "cancelled"}).json()["total"] == 1
    s = api.get(f"{API}/jobs/summary").json()
    assert s["accounts"][0]["email"] == "one@gmail.com" and "extraction_pending" in s
    assert api.post(f"{API}/jobs/{r.json()['id']}/retry").status_code == 409  # completed jobs are not retryable


def test_running_job_cancel_stops_at_checkpoint(api, db, connected_account, mailbox):
    from tests.helpers import jd_message
    from app.tasks.runner import sync_account_impl
    from app.services import jobs

    a = connected_account("one@gmail.com")
    box = mailbox.for_email("one@gmail.com")
    box.page_size = 1
    for i in range(4):
        box.add(jd_message(f"c0de000{i}", subject=f"Req {i}"))
    j = jobs.create_job(db, org_id=a.organization_id, job_type="sync", account_id=a.id)
    db.commit()
    orig = box.list_message_ids

    def cancelling(query, token, max_results=100):
        out = orig(query, token, max_results)
        job = db.get(ProcessingJob, j.id)
        job.status = "cancel_requested"
        db.commit()
        return out

    box.list_message_ids = cancelling
    sync_account_impl(str(j.id))
    job = db.get(ProcessingJob, j.id)
    assert job.status == "cancelled"
    from app.models.email import EmailMessage

    assert len(db.scalars(select(EmailMessage)).all()) == 1  # stopped after the first page
    assert db.get(EmailAccount, a.id).sync_state.gmail_history_id is None  # cursor is not advanced by a partial sync


def test_settings_validation_permissions_and_ai_status(api, client):
    g = api.get(f"{API}/settings/dedup").json()
    assert g["probable_threshold"] == 0.72
    assert api.put(f"{API}/settings/dedup", json={"values": {"probable_threshold": 0.9}}).json()["probable_threshold"] == 0.9
    assert api.put(f"{API}/settings/dedup", json={"values": {"probable_threshold": 5}}).status_code == 422
    assert api.get(f"{API}/settings/unknown").status_code == 404
    ai = api.get(f"{API}/settings-meta/ai").json()
    assert ai["configured"] is False and "api_key" not in str(ai).lower() and "paused" in ai["message"]


def test_analytics_reflect_real_data_and_empty_state(api):
    empty = api.get(f"{API}/analytics/overview").json()
    assert empty["kpis"]["total_unique_requirements"] == 0 and empty["top_skills"] == [] and empty["by_status"] == []
    rec = api.get(f"{API}/analytics/recruitment").json()
    assert rec["duplicate_rate"]["value"] is None and rec["extraction_success_rate"]["value"] is None  # no invented metrics
    now = utcnow().isoformat()
    a = mk(api, received_at=now, client_name="Globex", location="Pune")["id"]
    mk(api, received_at=now, client_name="Globex", required_skills=["Python", "AWS"], location="Pune", job_title="B")
    mk(api, received_at=(utcnow() - timedelta(days=3)).isoformat(), client_name="Initech", job_title="C", location="Delhi")
    api.post(f"{API}/requirements/{a}/approve")
    o = api.get(f"{API}/analytics/overview").json()
    k = o["kpis"]
    assert k["total_unique_requirements"] == 3 and k["approved"] == 1 and k["awaiting_review"] == 2 and k["received_today"] == 2
    assert {x["name"]: x["count"] for x in o["by_client"]} == {"Globex": 2, "Initech": 1}
    assert o["top_skills"][0]["name"] in ("Python",) and o["top_skills"][0]["count"] == 3
    assert sum(p["count"] for p in o["received_over_time"]) == 3
    rec = api.get(f"{API}/analytics/recruitment").json()
    assert rec["manual_correction_rate"]["value"] == 0 or rec["manual_correction_rate"]["value"] is None


def test_audit_log_records_user_actions_without_secrets(api):
    rid = mk(api)["id"]
    api.patch(f"{API}/requirements/{rid}", json={"job_title": "New"})
    api.post(f"{API}/requirements/{rid}/approve")
    logs = api.get(f"{API}/audit-logs", params={"entity_type": "requirement"}).json()["items"]
    assert {"requirement.created", "requirement.edited", "requirement.approved"} <= {x["action"] for x in logs}
    assert all(x["user_email"] == "owner@example.com" for x in logs)
    assert "password" not in str(api.get(f"{API}/audit-logs").json()).lower().replace("password_changed", "")


def test_health_endpoints(client, api):
    h = client.get(f"{API}/health")
    assert h.status_code == 200 and h.json()["status"] == "ok" and set(h.json()) == {"status", "database", "version"}
    d = api.get(f"{API}/health/details").json()
    assert d["database"]["status"] == "ok" and d["ai"]["status"] == "not_configured" and d["gmail"]["status"] == "ok" and d["worker"]["status"] == "not_configured"
    assert "secret" not in str(d).lower()
    anon = client.__class__(client.app)
    assert anon.get(f"{API}/health/details").status_code == 401


def test_archived_requirements_are_hidden_from_the_default_list_and_exports(api):
    keep, gone = mk(api, job_title="Keep")["id"], mk(api, job_title="Archive me")["id"]
    api.post(f"{API}/requirements/{gone}/archive")
    g = lambda **p: api.get(f"{API}/requirements", params=p).json()  # noqa: E731
    assert [i["job_title"] for i in g()["items"]] == ["Keep"]
    assert [i["job_title"] for i in g(lifecycle_status="archived")["items"]] == ["Archive me"]
    assert make_export(api, scope="filtered")["record_count"] == 1
    assert api.get(f"{API}/requirements/{keep}").status_code == 200 and api.get(f"{API}/requirements/{gone}").status_code == 200  # still reachable by id


def test_approving_clears_automatic_warnings(api, db):
    from app.models.requirements import JobRequirement

    rid = mk(api)["id"]
    row = db.scalar(select(JobRequirement).where(JobRequirement.id == uuid.UUID(rid)))
    row.validation_warnings = ["low_confidence", "conflict:salary_min", "job_title_not_found_in_source"]
    db.commit()
    assert len(api.get(f"{API}/requirements/{rid}").json()["validation_warnings"]) == 3
    assert api.post(f"{API}/requirements/{rid}/approve").json()["validation_warnings"] == []

