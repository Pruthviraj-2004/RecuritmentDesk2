import io

import pytest
from openpyxl import Workbook, load_workbook

from app.services.candidates import importer
from app.services.candidates.service import normalize_phone
from tests.conftest import login, make_user

API = "/api/v1"
MIN = {"first_name": "Asha", "email": "asha@example.com", "phone": "+91 98765 43210"}


def mk(api, **kw):
    r = api.post(f"{API}/candidates", json={**MIN, **kw})
    assert r.status_code == 201, r.text
    return r.json()


def upload(api, url, name, data: bytes, on_duplicate="skip", mime="text/csv"):
    return api.c.post(url, files={"file": (name, data, mime)}, data={"on_duplicate": on_duplicate}, headers=api._h())


def csv_bytes(*rows, header="first_name,last_name,email,phone,skills,total_experience_years,status,notes"):
    return ("\n".join([header, *rows]) + "\n").encode()


# ---------------------------------------------------------------- mandatory fields, validation, duplicates
def test_name_email_and_phone_are_mandatory_everything_else_optional(api):
    c = mk(api)  # only the three mandatory fields
    assert c["status"] == "new" and c["skills"] == [] and c["current_title"] is None and c["total_experience"] is None and c["source"] == "manual"
    assert c["full_name"] == "Asha"
    for missing in ("first_name", "email", "phone"):
        body = {k: v for k, v in MIN.items() if k != missing}
        r = api.post(f"{API}/candidates", json=body)
        assert r.status_code == 422, missing
    assert api.post(f"{API}/candidates", json={**MIN, "email": "b@example.com", "phone": "12"}).status_code == 422
    assert api.post(f"{API}/candidates", json={**MIN, "email": "not-an-email"}).status_code == 422
    assert api.post(f"{API}/candidates", json={**MIN, "email": "   "}).status_code == 422


def test_mandatory_fields_cannot_be_cleared_on_edit(api):
    c = mk(api)
    for field in ("first_name", "email", "phone"):
        r = api.patch(f"{API}/candidates/{c['id']}", json={field: ""})
        assert r.status_code in (422, 400), field
    ok = api.patch(f"{API}/candidates/{c['id']}", json={"phone": "9000000001", "current_title": "Dev", "rating": 4})
    assert ok.status_code == 200 and ok.json()["phone"] == "9000000001" and ok.json()["rating"] == 4


def test_field_validation(api):
    for bad in ({"total_experience": 90}, {"rating": 9}, {"status": "weird"}, {"expected_ctc": -5}, {"linkedin_url": "javascript:alert(1)"}, {"resume_url": "ftp://x"}):
        assert api.post(f"{API}/candidates", json={**MIN, "email": "z@example.com", "phone": "9111111111", **bad}).status_code == 422, bad


def test_duplicates_by_email_and_by_phone_in_any_format(api):
    a = mk(api)
    r = api.post(f"{API}/candidates", json={**MIN, "first_name": "Other", "email": "ASHA@example.com", "phone": "9000000000"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "duplicate_candidate" and r.json()["error"]["details"]["existing_id"] == a["id"]
    r = api.post(f"{API}/candidates", json={"first_name": "Other", "email": "other@example.com", "phone": "098765-43210"})
    assert r.status_code == 409  # same number, different formatting
    assert normalize_phone("+91 (987) 654-3210") == "9876543210" and normalize_phone("12") is None
    b = mk(api, email="b@example.com", phone="9222222222")
    assert api.patch(f"{API}/candidates/{b['id']}", json={"email": "asha@example.com"}).status_code == 409


# ---------------------------------------------------------------- list, filters, bulk, notes
def test_list_search_filters_sort_pagination_and_archive(api):
    mk(api, first_name="Asha", last_name="Verma", skills=["Python", "Django"], total_experience=6, current_location="Bengaluru", tags=["priority"], status="shortlisted", rating=5, notice_period="30 days")
    mk(api, first_name="Ravi", email="r@example.com", phone="9333333333", skills=["Java"], total_experience=2, current_location="Pune", status="new", notice_period="Immediate")
    for i in range(3):
        mk(api, first_name=f"Test{i}", email=f"t{i}@example.com", phone=f"900000000{i}", skills=["SQL"], total_experience=10 + i)
    g = lambda **p: api.get(f"{API}/candidates", params=p).json()  # noqa: E731
    assert g()["total"] == 5
    assert g(q="verma")["total"] == 1 and g(q="r@example")["total"] == 1 and g(q="django")["total"] == 1 and g(q="priority")["total"] == 1 and g(q="98765")["total"] == 1
    assert g(skill="python")["total"] == 1 and g(status="shortlisted")["total"] == 1 and g(location="pune")["total"] == 1 and g(tag="priority")["total"] == 1
    assert g(exp_min=9)["total"] == 3 and g(exp_max=3)["total"] == 1 and g(rating_min=5)["total"] == 1 and g(notice="immed")["total"] == 1
    assert g(q="100%")["total"] == 0
    p = g(page_size=2, sort="total_experience", direction="desc")
    assert [i["first_name"] for i in p["items"]] == ["Test2", "Test1"] and p["total"] == 5
    cid = g(q="ravi")["items"][0]["id"]
    assert api.delete(f"{API}/candidates/{cid}").json() == {"archived": True}
    assert g()["total"] == 4 and g(archived="true")["total"] == 1
    assert api.post(f"{API}/candidates/{cid}/restore").json()["is_archived"] is False and g()["total"] == 5
    assert api.get(f"{API}/candidates", params={"page_size": 999}).status_code == 422


def test_stats_bulk_actions_and_notes_timeline(api):
    ids = [mk(api, email=f"c{i}@example.com", phone=f"91000000{i:02d}")["id"] for i in range(3)]
    r = api.post(f"{API}/candidates/bulk", json={"ids": ids + ["00000000-0000-0000-0000-000000000001"], "action": "set_status", "value": "screening"}).json()
    assert len(r["succeeded"]) == 3 and len(r["failed"]) == 1
    assert api.post(f"{API}/candidates/bulk", json={"ids": ids, "action": "set_status", "value": "bogus"}).status_code == 422
    api.post(f"{API}/candidates/bulk", json={"ids": ids[:2], "action": "add_tag", "value": "Hot"})
    api.post(f"{API}/candidates/bulk", json={"ids": ids[:1], "action": "remove_tag", "value": "hot"})
    api.post(f"{API}/candidates/bulk", json={"ids": ids[2:], "action": "archive"})
    assert api.get(f"{API}/candidates/{ids[1]}").json()["tags"] == ["Hot"] and api.get(f"{API}/candidates/{ids[0]}").json()["tags"] == []
    s = api.get(f"{API}/candidates/stats").json()
    assert s["total"] == 2 and s["by_status"] == {"screening": 2} and s["new_this_week"] == 2
    n = api.post(f"{API}/candidates/{ids[0]}/notes", json={"body": "Called, interested in Pune role"})
    assert n.status_code == 201 and n.json()["author"] == "Owner"
    tl = api.get(f"{API}/candidates/{ids[0]}/notes").json()
    assert [x["kind"] for x in tl] == ["note", "status"] and "screening" in tl[1]["body"]  # newest first; status changes are logged
    assert api.get(f"{API}/candidates/{ids[0]}").json()["last_contacted_at"] is not None


# ---------------------------------------------------------------- import
def test_template_downloads_and_example_rows_import_cleanly(api):
    for fmt in ("csv", "xlsx"):
        r = api.get(f"{API}/candidates/import/template", params={"format": fmt})
        assert r.status_code == 200 and f"candidate-import-template.{fmt}" in r.headers["content-disposition"]
        prev = upload(api, f"{API}/candidates/import/preview", f"t.{fmt}", r.content)
        assert prev.status_code == 200, prev.text
        j = prev.json()
        assert j["counts"]["create"] == 2 and j["counts"]["error"] == 0 and len(j["recognised_columns"]) == len(importer.STANDARD_HEADERS)
    wb = load_workbook(io.BytesIO(api.get(f"{API}/candidates/import/template").content))
    assert wb.sheetnames == ["Candidates", "Instructions"] and [c.value for c in wb["Candidates"][1]] == importer.STANDARD_HEADERS
    assert api.get(f"{API}/candidates/import/template", params={"format": "pdf"}).status_code == 422


def test_preview_saves_nothing_then_commit_saves(api):
    data = csv_bytes("Asha,Verma,asha@example.com,9876543210,Python;Django,6,shortlisted,Spoke on Monday", "Ravi,,ravi@example.com,9333333333,Java,2,new,")
    p = upload(api, f"{API}/candidates/import/preview", "c.csv", data).json()
    assert p["counts"]["create"] == 2 and api.get(f"{API}/candidates").json()["total"] == 0  # dry run
    r = upload(api, f"{API}/candidates/import", "c.csv", data)
    assert r.status_code == 201 and r.json()["created"] == 2 and r.json()["failed"] == 0
    got = api.get(f"{API}/candidates", params={"q": "asha"}).json()["items"][0]
    assert got["skills"] == ["Django", "Python"] and got["total_experience"] == 6 and got["status"] == "shortlisted" and got["source"] == "import"
    assert api.get(f"{API}/candidates/{got['id']}/notes").json()[0]["body"] == "Spoke on Monday"
    h = api.get(f"{API}/candidates/imports").json()
    assert h[0]["filename"] == "c.csv" and h[0]["created"] == 2


def test_mandatory_fields_are_enforced_per_row_with_clear_messages(api):
    data = csv_bytes(",X,x@example.com,9876543210,Python,1,new,", "NoEmail,Y,,9876543211,Python,1,new,", "NoPhone,Z,z@example.com,,Python,1,new,", "Bad,Q,notmail,12,Python,99,wrong,", "Good,G,good@example.com,9999999999,,,,")
    j = upload(api, f"{API}/candidates/import/preview", "c.csv", data).json()
    assert j["counts"]["create"] == 1 and j["counts"]["error"] == 4
    by_row = {p["row"]: p["errors"] for p in j["problems"]}
    assert "name is missing" in by_row[2] and "email is required" in by_row[3] and "phone is required" in by_row[4]
    assert any("email" in e for e in by_row[5]) and any("phone" in e for e in by_row[5]) and any("between 0 and 60" in e for e in by_row[5])
    r = upload(api, f"{API}/candidates/import", "c.csv", data).json()
    assert r["created"] == 1 and r["failed"] == 4 and api.get(f"{API}/candidates").json()["total"] == 1  # good rows still import


def test_flexible_headers_full_name_and_semicolon_delimiter(api):
    data = b"Candidate Name;Email ID;Mobile No;Key Skills;Experience;CTC;ECTC;Location;Notice\nAsha Rao Verma;asha@example.com;98765 43210;Python, Django | AWS;5 years;12 LPA;18,00,000;Pune;Immediate\n"
    j = upload(api, f"{API}/candidates/import", "odd.csv", data)
    assert j.status_code == 201 and j.json()["created"] == 1, j.text
    c = api.get(f"{API}/candidates").json()["items"][0]
    assert (c["first_name"], c["last_name"]) == ("Asha", "Rao Verma") and c["skills"] == ["AWS", "Django", "Python"] and c["total_experience"] == 5
    assert c["current_ctc"] == 12 and c["expected_ctc"] == 1800000 and c["current_location"] == "Pune" and c["notice_period"] == "Immediate"


def test_duplicates_skip_or_update_and_in_file_duplicates(api):
    mk(api, current_title="Old title", skills=["Python"], tags=["a"])
    data = csv_bytes("Asha,V,asha@example.com,9000000000,Django,7,new,", "Dup,Row,other@example.com,98765 43210,Go,,,", "Fresh,F,fresh@example.com,9555555555,Rust,1,,", "Fresh2,F,FRESH@example.com,9555555556,Rust,1,,")
    skip = upload(api, f"{API}/candidates/import/preview", "d.csv", data, "skip").json()
    assert skip["counts"] == {"create": 1, "update": 0, "skip_duplicate": 2, "skip_in_file": 1, "error": 0}
    assert any(p["matched_name"] == "Asha" for p in skip["problems"])
    upd = upload(api, f"{API}/candidates/import", "d.csv", data, "update").json()
    assert upd["created"] == 1 and upd["updated"] == 2 and upd["skipped"] == 1
    c = api.get(f"{API}/candidates", params={"q": "asha@example"}).json()["items"][0]
    assert c["current_title"] == "Old title" and c["skills"] == ["Django", "Go", "Python"] and c["total_experience"] == 7  # filled/merged, nothing blanked
    assert api.get(f"{API}/candidates").json()["total"] == 2
    again = upload(api, f"{API}/candidates/import", "d.csv", data, "skip")
    assert again.status_code == 422 and again.json()["error"]["code"] == "nothing_to_import"


def test_import_rejects_bad_files(api):
    url = f"{API}/candidates/import/preview"
    assert upload(api, url, "e.csv", b"").status_code == 422
    assert upload(api, url, "no_name.csv", b"email,phone\na@b.com,9876543210\n").json()["error"]["code"] == "invalid_template"
    assert upload(api, url, "no_contact.csv", b"first_name,skills\nA,Python\n").status_code == 422
    assert upload(api, url, "old.xls", b"\xd0\xcf\x11\xe0junk").status_code == 422
    assert upload(api, url, "big.csv", b"a,b\n" + b"x" * (importer.MAX_BYTES + 10)).status_code == 422
    assert upload(api, url, "dup.csv", csv_bytes("A,B,a@b.com,9876543210,,,,"), on_duplicate="merge").status_code == 422
    many = csv_bytes(*[f"N{i},L,n{i}@example.com,9{i:09d},,,," for i in range(importer.MAX_ROWS + 5)])
    assert upload(api, url, "many.csv", many).status_code == 422
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", "<x/>")
        z.writestr("xl/vbaProject.bin", b"macro")
    assert upload(api, url, "macro.xlsx", buf.getvalue()).status_code == 422
    assert upload(api, url, "notexcel.xlsx", b"PK\x03\x04garbage").status_code == 422


def test_xlsx_upload_with_numbers_and_formulas_as_plain_values(api):
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    ws.append(["First Name", "Email", "Phone", "Experience", "Skills"])
    ws.append(["@SUM(A1)", "x@example.com", 9876543210, 4.5, "Python"])  # text that merely looks like a formula
    ws.append(["Real", "y@example.com", 9876543211, None, None])
    ws["A4"] = "=1+1"  # a real formula cell: never evaluated, so the (name-less) row is rejected, not executed
    ws["B4"], ws["C4"] = "z@example.com", 9876543212
    out = io.BytesIO()
    wb.save(out)
    r = upload(api, f"{API}/candidates/import", "c.xlsx", out.getvalue(), mime="application/octet-stream")
    assert r.status_code == 201 and r.json()["created"] == 2 and r.json()["failed"] == 1
    c = api.get(f"{API}/candidates", params={"q": "x@example"}).json()["items"][0]
    assert c["phone"] == "9876543210" and c["total_experience"] == 4.5
    exp = api.post(f"{API}/candidates/export", json={})
    assert exp.status_code == 200
    rows = {r[0].value: r[0] for r in load_workbook(io.BytesIO(exp.content)).worksheets[0].iter_rows(min_row=2)}
    cell = next(c for k, c in rows.items() if "SUM" in str(k))
    assert cell.data_type != "f" and str(cell.value).startswith("'@")  # formula-looking text is neutralised on export


# ---------------------------------------------------------------- export
def test_export_selected_and_filtered(api):
    a = mk(api, skills=["Python"], status="shortlisted")
    mk(api, first_name="Ravi", email="r@example.com", phone="9333333333")
    e = api.post(f"{API}/candidates/export", json={"ids": [a["id"]]})
    ws = load_workbook(io.BytesIO(e.content)).worksheets[0]
    assert "spreadsheetml" in e.headers["content-type"] and ws.max_row == 2 and ws["A2"].value == "Asha" and ws["H2"].value == "Python" and ws.freeze_panes == "A2"
    assert load_workbook(io.BytesIO(api.post(f"{API}/candidates/export", json={"filters": {"status": "shortlisted"}}).content)).worksheets[0].max_row == 2
    assert api.post(f"{API}/candidates/export", json={"filters": {"status": "hired"}}).status_code == 422


# ---------------------------------------------------------------- pipeline & matching
def req(api, **kw):
    body = {"job_title": "Senior Python Developer", "client_name": "Globex", "required_skills": ["Python", "Django", "AWS"], "minimum_experience": 5, "maximum_experience": 8,
            "location": "Bengaluru, India", "work_mode": "onsite", **kw}
    r = api.post(f"{API}/requirements", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def test_matching_scores_are_explained_and_ranked(api):
    r = req(api)
    best = mk(api, first_name="Best", email="best@example.com", phone="9100000001", skills=["Python", "Django", "AWS"], total_experience=6, current_location="Bengaluru")
    mid = mk(api, first_name="Mid", email="mid@example.com", phone="9100000002", skills=["Python", "Flask"], total_experience=9, current_location="Pune", willing_to_relocate=True)
    mk(api, first_name="None", email="none@example.com", phone="9100000003", skills=["Cooking"], total_experience=6)
    rejected = mk(api, first_name="Rej", email="rej@example.com", phone="9100000004", skills=["Python", "Django", "AWS"], status="rejected")
    m = api.get(f"{API}/requirements/{r['id']}/matching-candidates").json()
    names = [x["candidate"]["first_name"] for x in m]
    assert names[0] == "Best" and "Mid" in names and "None" not in names and "Rej" not in names  # rejected / unrelated excluded
    assert m[0]["score"] >= 95 and m[0]["missing_skills"] == [] and any("3 of 3" in x for x in m[0]["reasons"])
    mid_m = next(x for x in m if x["candidate"]["first_name"] == "Mid")
    assert mid_m["score"] < m[0]["score"] and mid_m["missing_skills"] == ["AWS", "Django"] or set(mid_m["missing_skills"]) == {"Django", "AWS"}
    back = api.get(f"{API}/candidates/{best['id']}/matching-requirements").json()
    assert back[0]["requirement"]["id"] == r["id"] and back[0]["score"] == m[0]["score"]
    assert rejected["id"]


def test_pipeline_add_update_remove_and_no_duplicates(api):
    r = req(api)
    c = mk(api, skills=["Python", "Django"])
    s = api.post(f"{API}/candidates/{c['id']}/submissions", json={"requirement_id": r["id"], "stage": "submitted", "notes": "sent to client"})
    assert s.status_code == 201 and s.json()["stage"] == "submitted" and s.json()["requirement_title"] == "Senior Python Developer"
    again = api.post(f"{API}/candidates/{c['id']}/submissions", json={"requirement_id": r["id"]})
    assert again.status_code == 409 and again.json()["error"]["code"] == "already_in_pipeline"
    sid = s.json()["id"]
    assert api.patch(f"{API}/candidates/submissions/{sid}", json={"stage": "interview"}).json()["stage"] == "interview"
    assert api.patch(f"{API}/candidates/submissions/{sid}", json={"stage": "nope"}).status_code == 422
    assert api.get(f"{API}/requirements/{r['id']}/submissions").json()[0]["candidate_name"] == "Asha"
    assert api.get(f"{API}/candidates/{c['id']}").json()["pipeline_count"] == 1
    assert api.get(f"{API}/candidates/stats").json()["in_pipeline"] == 1
    tl = " ".join(n["body"] for n in api.get(f"{API}/candidates/{c['id']}/notes").json())
    assert "Added to pipeline" in tl and "changed from submitted to interview" in tl
    m = api.get(f"{API}/requirements/{r['id']}/matching-candidates").json()
    assert m and m[0]["stage"] == "interview"  # matches show who is already in the pipeline
    assert api.delete(f"{API}/candidates/submissions/{sid}").json() == {"removed": True}
    assert api.post(f"{API}/candidates/{c['id']}/submissions", json={"requirement_id": "00000000-0000-0000-0000-000000000001"}).status_code == 404


# ---------------------------------------------------------------- security
def test_candidates_are_isolated_between_organizations_and_viewers_cannot_write(client, api, db):
    from sqlalchemy import select

    from app.models.auth import OrganizationMembership
    from app.models.enums import Role

    c = mk(api)
    make_user(db, "other@example.com", org="Org2")
    o = login(client.__class__(client.app), "other@example.com")
    assert o.get(f"{API}/candidates").json()["total"] == 0
    assert o.get(f"{API}/candidates/{c['id']}").status_code == 404
    assert o.patch(f"{API}/candidates/{c['id']}", json={"rating": 1}).status_code == 404
    assert o.get(f"{API}/candidates/{c['id']}/notes").status_code == 404
    assert o.post(f"{API}/candidates/export", json={"ids": [c["id"]]}).status_code == 422  # nothing from another org is exported
    # same email in another org is allowed (uniqueness is per organization)
    assert o.post(f"{API}/candidates", json=MIN).status_code == 201
    viewer_user, _ = make_user(db, "viewer@example.com", org="Tmp")
    m = db.scalar(select(OrganizationMembership).where(OrganizationMembership.user_id == viewer_user.id))
    owner_org = db.scalar(select(OrganizationMembership).join(OrganizationMembership.user).where(OrganizationMembership.user.has(email="owner@example.com"))).organization_id
    m.organization_id, m.role = owner_org, Role.VIEWER
    db.commit()
    v = login(client.__class__(client.app), "viewer@example.com")
    assert v.get(f"{API}/candidates").json()["total"] == 1
    assert v.post(f"{API}/candidates", json={**MIN, "email": "v@example.com", "phone": "9777777777"}).status_code == 403
    assert v.post(f"{API}/candidates/import", data={"on_duplicate": "skip"}, files={"file": ("a.csv", csv_bytes("A,B,a@b.com,9876543210,,,,"), "text/csv")}, headers=v._h()).status_code == 403
    assert v.post(f"{API}/candidates/bulk", json={"ids": [c["id"]], "action": "archive"}).status_code == 403
    anon = client.__class__(client.app)
    assert anon.get(f"{API}/candidates").status_code == 401


# ---------------------------------------------------------------- demo data + phrase-style skills
def test_matching_understands_phrase_style_requirement_skills():
    from app.services.candidates.matching import skill_matches

    c = {"docker": "Docker", "sql": "SQL", "java": "Java", "go": "Go", "python": "Python"}
    assert skill_matches("docker, kubernetes, and cloud platforms", c) and skill_matches("advanced sql queries and joins", c)
    assert skill_matches("core java and object-oriented programming", c) and skill_matches("python", c)
    assert not skill_matches("javascript frameworks", c)  # Java is not JavaScript
    assert not skill_matches("good to go", c)  # 1-2 letter skills only match exactly
    assert not skill_matches("terraform", c)


def test_demo_candidates_are_marked_varied_idempotent_and_removable(api, owner):
    from app.models.candidates import Candidate, CandidateSubmission
    from app.services.candidates import demo
    from sqlalchemy import func, select

    from app.db.session import session_factory

    org_id, uid = owner[1].id, owner[0].id
    db = session_factory()()  # a plain session (the shared test session expires rows while eager-loading)
    r1 = req(api, job_title="Cloud DevOps Engineer", required_skills=["Docker and Kubernetes", "AWS services: EC2, S3, IAM", "CI/CD pipeline implementation", "Terraform", "Linux"], minimum_experience=4, maximum_experience=6, location="Pune")
    req(api, job_title="Data Analyst", required_skills=["Advanced SQL queries and joins", "Basic Python with Pandas", "Power BI dashboards"], minimum_experience=1, maximum_experience=3, location="Hyderabad")
    plan = demo.seed(db, org_id, uid, apply=False)
    assert plan["requirements"] == 2 and 6 <= plan["created"] <= 10
    assert db.scalar(select(func.count()).select_from(Candidate)) == 0  # a dry run writes nothing
    done = demo.seed(db, org_id, uid, apply=True)
    db.commit()
    assert done["created"] == plan["created"]
    cands = list(db.scalars(select(Candidate)))
    assert all(c.source == "demo" and "demo" in c.tags and c.email.endswith("@demo-candidates.example.org") and c.phone.startswith("90000") for c in cands)
    assert len({c.email for c in cands}) == len(cands) and len({c.phone_normalized for c in cands}) == len(cands)
    m = api.get(f"{API}/requirements/{r1['id']}/matching-candidates").json()
    scores = [x["score"] for x in m]
    assert m and scores == sorted(scores, reverse=True) and scores[0] >= 85  # a strong fit exists
    assert len(set(scores)) >= 3 and min(scores) < scores[0] - 20  # and the spread shows different match qualities
    assert db.scalar(select(func.count()).select_from(CandidateSubmission)) >= 1  # at least one demo pipeline entry
    again = demo.seed(db, org_id, uid, apply=True)
    assert again["created"] == 0 and again["skipped_existing"] == plan["created"]  # re-running never duplicates
    db.commit()
    real = mk(api, first_name="Real", email="real@client.com", phone="9444444444")
    db.rollback()  # drop the read snapshot taken before the API wrote
    assert demo.purge(db, apply=False) == plan["created"]
    assert demo.purge(db, apply=True) == plan["created"]
    db.commit()
    left = db.scalars(select(Candidate)).all()
    assert [c.email for c in left] == [real["email"]]  # real candidates are never touched
    db.close()


def test_blank_template_has_headers_only_and_cannot_be_imported_by_accident(api):
    r = api.get(f"{API}/candidates/import/template", params={"format": "xlsx", "examples": "false"})
    assert "candidate-import-template-blank.xlsx" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content))["Candidates"]
    assert [c.value for c in ws[1]] == importer.STANDARD_HEADERS and ws.max_row == 1  # headers only
    csv_blank = api.get(f"{API}/candidates/import/template", params={"format": "csv", "examples": "false"}).content
    assert csv_blank.decode("utf-8-sig").strip().splitlines() == [",".join(importer.STANDARD_HEADERS)]
    j = upload(api, f"{API}/candidates/import", "blank.xlsx", r.content, mime="application/octet-stream")
    assert j.status_code == 422 and j.json()["error"]["code"] == "nothing_to_import"
    assert api.get(f"{API}/candidates").json()["total"] == 0

