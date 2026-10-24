import io
import zipfile

import pytest
from openpyxl import load_workbook

from app.core.exceptions import ValidationFailure
from app.core.logging import redact
from app.core.security import TokenCipher, hash_password, validate_password_strength, verify_password
from app.models.requirements import JobRequirement
from app.services.deduplication.matcher import Profile, compare, decide, ordered_pair
from app.services.delivery.service import build_mime, clean_header, validate_recipients
from app.services.exports.excel import COLUMNS, build_workbook, resolve_columns, safe_cell, safe_filename, safe_sheet_name
from app.services.gmail.attachments import parse_attachment
from app.services.extraction.schema import ExtractedRequirement
from app.services.extraction.validation import validate_against_source


def prof(**kw):
    base = dict(client="globex", title="senior python developer", ref="", required_skills=frozenset({"python", "django", "aws"}),
                exp_min=5, exp_max=8, location="bangalore")
    base.update(kw)
    return Profile(**base)


# ---------------------------------------------------------------- matcher
def test_same_client_and_reference_is_exact():
    d = decide(prof(ref="gx1042"), prof(ref="gx1042", title="python lead"))
    assert d.kind == "exact"


def test_different_reference_ids_never_merge():
    d = decide(prof(ref="gx1"), prof(ref="gx2"))
    assert d.kind in ("related", "distinct") and d.kind != "exact"
    assert decide(prof(ref="gx1"), prof(ref="gx2", title="data engineer")).kind == "distinct"


def test_title_match_alone_is_not_a_duplicate():
    a = Profile(title="python developer")
    b = Profile(title="python developer")
    assert decide(a, b).kind == "insufficient"
    a2 = prof(client="alpha", required_skills=frozenset({"java"}), location="delhi", exp_min=1, exp_max=2)
    b2 = prof(client="beta", required_skills=frozenset({"python"}), location="pune", exp_min=10, exp_max=12)
    assert decide(a2, b2).kind == "distinct"


def test_probable_duplicate_when_many_signals_match_without_ids():
    a = prof(recruiter_email="j@x.com", description="x" * 100 + " python django developer needed")
    b = prof(title="python developer", recruiter_email="j@x.com", location="bangalore")
    d = decide(a, b)
    assert d.kind in ("exact", "probable") and d.score > 0.7


def test_same_role_far_apart_in_time_is_not_auto_matched():
    from datetime import datetime, timedelta, timezone

    t = datetime(2026, 1, 1, tzinfo=timezone.utc)
    a = prof(received_at=t, fingerprint="f")
    b = prof(received_at=t + timedelta(days=400), fingerprint="f")
    assert decide(a, b).kind in ("probable", "distinct")


def test_conflicts_reported():
    c = compare(prof(salary_min=100.0, work_mode="remote"), prof(salary_min=120.0, work_mode="onsite"))
    assert {"salary_min", "work_mode"} <= set(c.conflicts)


def test_pair_ordering_is_canonical():
    assert ordered_pair("b", "a") == ordered_pair("a", "b") == ("a", "b")


# ---------------------------------------------------------------- extraction validation
def test_ai_output_is_validated_against_source():
    src = "Client: Globex. Python Developer. Skills: Python, Django. Contact jane@staffing.com"
    ex = ExtractedRequirement(client_name="Globex", job_title="Python Developer", required_skills=["Python", "Django", "Rust"],
                              recruiter_email="invented@evil.com", evidence={"client_name": "Client: Globex", "job_title": "made up snippet"},
                              minimum_experience=9, maximum_experience=3)
    cleaned, warnings, conf = validate_against_source(ex, src)
    assert "Rust" not in cleaned.required_skills
    assert cleaned.recruiter_email is None and not any("recruiter_email" in w for w in warnings)  # dropped silently: nothing left to verify
    assert cleaned.evidence == {"client_name": "Client: Globex"}
    assert (cleaned.minimum_experience, cleaned.maximum_experience) == (3, 9) and "experience_range_inverted" in warnings
    assert 0 < conf < 1


def test_schema_rejects_bad_values():
    with pytest.raises(Exception):
        ExtractedRequirement(minimum_experience=-1)
    assert ExtractedRequirement(work_mode="On-Site").work_mode == "onsite"
    assert ExtractedRequirement(work_mode="whatever").work_mode == "unknown"
    assert ExtractedRequirement(required_skills="a, b;c").required_skills == ["a", "b", "c"]


# ---------------------------------------------------------------- attachments
def _docx_bytes(text="Python developer job"):
    import docx

    d = docx.Document()
    d.add_paragraph(text)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_attachment_type_is_sniffed_not_trusted():
    ok = parse_attachment("jd.docx", "application/octet-stream", _docx_bytes(), 1_000_000)
    assert ok.status == "extracted" and "Python developer" in ok.text
    exe = parse_attachment("resume.pdf", "application/pdf", b"MZ\x90\x00" + b"\x00" * 100, 1_000_000)
    assert exe.status == "quarantined"
    assert parse_attachment("big.txt", "text/plain", b"a" * 2000, 1000).status == "too_large"
    assert parse_attachment("x.bin", None, bytes(range(256)) * 4, 10_000).status == "unsupported"
    assert parse_attachment("a.txt", "text/plain", "hello job".encode(), 1000).text == "hello job"


def test_macro_docs_and_zip_bombs_are_quarantined():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<x/>")
        z.writestr("word/vbaProject.bin", b"macro")
    assert parse_attachment("m.docx", None, buf.getvalue(), 10**7).status == "quarantined"
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * 5_000_000)
    assert parse_attachment("b.docx", None, bomb.getvalue(), 10**7).status == "quarantined"
    assert parse_attachment("p.pdf", None, b"%PDF-1.4 garbage", 10**6).status == "failed"


# ---------------------------------------------------------------- excel
def test_formula_injection_is_neutralised():
    for v in ("=cmd|' /C calc'!A0", "+1+1", "-2+3", "@SUM(A1)", "\tcmd"):
        assert safe_cell(v).startswith("'")
    assert safe_cell("normal") == "normal" and safe_cell(5) == 5


def test_sheet_and_filename_sanitization():
    assert safe_sheet_name("a/b:c*?[x]") == "a b c   x"
    assert len(safe_sheet_name("x" * 80)) == 31
    assert safe_filename("../../etc/passwd") == "etcpasswd"
    assert safe_filename("") == "recruitment-requirements"


def _req(i, **kw):
    import uuid

    r = JobRequirement(id=uuid.uuid4(), organization_id=uuid.uuid4(), job_title=f"Dev {i}", client_name="=HYPERLINK(\"x\")", work_mode="remote", review_status="approved",
                       lifecycle_status="active", certifications=[], required_documents=[], validation_warnings=[], version_number=1, **kw)
    return r


def test_workbook_structure_and_dedup_of_ids():
    reqs = [_req(1, location="Pune"), _req(2)]
    cols = resolve_columns(None, include_source=False, include_audit=True)
    data, n = build_workbook(reqs + [reqs[0]], cols, sheet_name="Weekly: Report", summary={}, metadata={"Filters": {"a": 1}})
    assert n == 2  # the repeated canonical id is exported once
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Weekly  Report", "Summary", "Export Info"]
    ws = wb.worksheets[0]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws.max_row == 3
    assert ws["A1"].value == "Client" and str(ws["A2"].value).startswith("'=")  # formula escaped
    headers = [c.value for c in ws[1]]
    assert "Review Status" in headers and "ID" not in headers  # audit columns opt-in, no internal ids


def test_column_selection_ignores_unknown_and_defaults():
    assert [c.key for c in resolve_columns(["job_title", "nope", "client_name", "job_title"], False, False)] == ["job_title", "client_name"]
    assert resolve_columns(["nope"], False, False)  # falls back to defaults
    assert len({c.key for c in COLUMNS}) == len(COLUMNS)


# ---------------------------------------------------------------- delivery validation & security
def test_recipient_validation():
    assert validate_recipients(["A@Example.com", "a@example.com", " b@example.org "]) == ["a@example.com", "b@example.org"]
    for bad in ("not-an-email", "a@b.com\r\nBcc: evil@x.com", "a@b.com, c@d.com", "<a@b.com>"):
        with pytest.raises(ValidationFailure):
            validate_recipients([bad])
    with pytest.raises(ValidationFailure):
        validate_recipients([f"u{i}@example.com" for i in range(25)])


def test_header_injection_blocked_and_mime_has_attachment():
    assert "\n" not in clean_header("Hi\r\nBcc: x@y.com")
    msg, rid = build_mime(sender="me@gmail.com", to=["a@example.com"], cc=[], subject="Report\r\nBcc: evil@x.com", body="hello", filename="r.xlsx", xlsx=b"PK\x03\x04data")
    assert msg["Bcc"] is None and rid.startswith("<")
    atts = [p for p in msg.iter_attachments()]
    assert len(atts) == 1 and atts[0].get_filename() == "r.xlsx" and atts[0].get_content_type().endswith("spreadsheetml.sheet")


def test_token_encryption_roundtrip_and_redaction():
    from cryptography.fernet import Fernet

    c = TokenCipher(Fernet.generate_key().decode())
    enc = c.encrypt("1//refresh-token-secret")
    assert "refresh-token-secret" not in enc and c.decrypt(enc) == "1//refresh-token-secret"
    text = redact("failed with access_token=ya29.abcDEF123 and Bearer abc.def.ghi password: hunter2 sk-abcdefghijklmnopqrstuv")
    for secret in ("ya29.abcDEF123", "hunter2", "abc.def.ghi", "sk-abcdefghijklmnopqrstuv"):
        assert secret not in text


def test_password_hashing_and_strength():
    h = hash_password("Str0ngPassword!23")
    assert verify_password("Str0ngPassword!23", h) and not verify_password("wrong", h) and not verify_password("x", None)
    assert validate_password_strength("Ab1") and validate_password_strength("abcdefg1") and not validate_password_strength("Abcdef12")  # 8 chars is enough


def test_reformatted_copy_of_same_requirement_is_an_exact_duplicate_but_other_cities_are_not():
    """A digest/forward drops the job ID and reformats text: same client + title + compatible location is the same job."""
    a = prof(location="bangalore india")
    assert decide(a, prof(location="")).kind == "exact"                    # location missing in the copy
    assert decide(a, prof(location="bangalore")).kind == "exact"           # contained
    assert decide(a, prof(title="senior python developer ")).kind == "exact"
    assert decide(a, prof(location="pune india")).kind != "exact"          # a different city is a different opening
    assert decide(a, prof(client="initech", location="")).kind != "exact"  # a different client never matches
    assert decide(prof(ref="gx1"), prof(ref="gx2", location="")).kind != "exact"  # contradicting job IDs never merge
