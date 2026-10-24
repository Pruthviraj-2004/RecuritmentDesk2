"""Excel workbook generation with openpyxl. One row per canonical requirement; untrusted values are formula-escaped."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.requirements import JobRequirement

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def safe_cell(value: Any) -> Any:
    """Neutralise spreadsheet formula injection: text that could be evaluated gets a leading apostrophe-style guard."""
    if isinstance(value, str):
        value = _ILLEGAL.sub("", value)
        if value.startswith(FORMULA_PREFIXES):
            return "'" + value
        return value[:32000]
    return value


def safe_sheet_name(name: str, default: str = "Requirements") -> str:
    name = re.sub(r"[\[\]\:\*\?\/\\]", " ", name or "").strip().strip("'")[:31]
    return name or default


def safe_filename(name: str, default: str = "recruitment-requirements") -> str:
    name = re.sub(r"[^A-Za-z0-9._\- ]", "", name or "").strip().replace(" ", "-")[:100].strip(".-")
    return name or default


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    width: int
    getter: Callable[[JobRequirement], Any]
    kind: str = "text"  # text | number | date | wrap
    group: str = "core"  # core | source | audit


def _salary(r: JobRequirement) -> str | None:
    if r.salary_min is None and r.salary_max is None:
        return None
    cur = r.salary_currency or ""
    lo = f"{float(r.salary_min):,.0f}" if r.salary_min is not None else ""
    hi = f"{float(r.salary_max):,.0f}" if r.salary_max is not None else ""
    rng = f"{lo} - {hi}" if lo and hi else (lo or hi)
    return f"{cur} {rng}".strip() + (f" / {r.salary_period}" if r.salary_period else "")


def _exp(r: JobRequirement) -> str | None:
    lo, hi = r.minimum_experience, r.maximum_experience
    f = lambda v: f"{v:g}"  # noqa: E731
    if lo is None and hi is None:
        return None
    if lo is not None and hi is not None and hi != lo:
        return f"{f(lo)}-{f(hi)} yrs"
    return f"{f(lo if lo is not None else hi)}+ yrs" if lo is not None else f"up to {f(hi)} yrs"


def _naive(dt: datetime | None) -> datetime | None:
    return dt.astimezone(timezone.utc).replace(tzinfo=None) if dt else None


COLUMNS: list[Column] = [
    Column("client_name", "Client", 24, lambda r: r.client_name),
    Column("job_title", "Job Title", 32, lambda r: r.job_title),
    Column("job_reference_id", "Job ID", 16, lambda r: r.job_reference_id),
    Column("number_of_openings", "Openings", 10, lambda r: r.number_of_openings, "number"),
    Column("required_skills", "Required Skills", 42, lambda r: ", ".join(r.required_skills) or None, "wrap"),
    Column("preferred_skills", "Preferred Skills", 36, lambda r: ", ".join(r.preferred_skills) or None, "wrap"),
    Column("experience", "Experience", 14, _exp),
    Column("location", "Location", 24, lambda r: r.location),
    Column("work_mode", "Work Mode", 12, lambda r: None if r.work_mode == "unknown" else r.work_mode),
    Column("employment_type", "Employment Type", 16, lambda r: r.employment_type),
    Column("salary", "Salary / Rate", 22, _salary),
    Column("notice_period", "Notice Period", 14, lambda r: r.notice_period),
    Column("recruiter_name", "Recruiter", 20, lambda r: r.recruiter_name),
    Column("recruiter_email", "Recruiter Email", 28, lambda r: r.recruiter_email),
    Column("recruiter_phone", "Recruiter Phone", 16, lambda r: r.recruiter_phone),
    Column("submission_email", "Submission Email", 28, lambda r: r.submission_email),
    Column("submission_instructions", "Submission Instructions", 40, lambda r: r.submission_instructions, "wrap"),
    Column("job_description", "Job Description", 60, lambda r: r.job_description, "wrap"),
    Column("received_at", "Received", 14, lambda r: _naive(r.received_at), "date"),
    Column("review_status", "Review Status", 14, lambda r: r.review_status, group="audit"),
    Column("lifecycle_status", "Lifecycle", 12, lambda r: r.lifecycle_status, group="audit"),
    Column("last_updated_at", "Last Updated", 16, lambda r: _naive(r.last_updated_at), "date", "audit"),
    Column("source_accounts", "Source Accounts", 30, lambda r: ", ".join(getattr(r, "_source_accounts", [])) or None, "wrap", "source"),
    Column("source_count", "Source Emails", 10, lambda r: len(r.sources), "number", "source"),
]
COLUMN_MAP = {c.key: c for c in COLUMNS}
DEFAULT_COLUMNS = [c.key for c in COLUMNS if c.group == "core" and c.key not in ("job_description", "submission_instructions", "recruiter_phone", "notice_period", "preferred_skills", "submission_email")]
SORTABLE = {"received_at", "client_name", "job_title", "last_updated_at", "review_status"}

HEADER_FILL = PatternFill("solid", start_color="C2410C")  # brand orange
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
THIN = Side(style="thin", color="D4D4D8")


def resolve_columns(keys: list[str] | None, include_source: bool, include_audit: bool) -> list[Column]:
    chosen = [k for k in (keys or DEFAULT_COLUMNS) if k in COLUMN_MAP]
    cols = [COLUMN_MAP[k] for k in dict.fromkeys(chosen)]
    if include_source:
        cols += [c for c in COLUMNS if c.group == "source" and c not in cols]
    if include_audit:
        cols += [c for c in COLUMNS if c.group == "audit" and c not in cols]
    return cols or [COLUMN_MAP[k] for k in DEFAULT_COLUMNS]


def build_workbook(rows: Iterable[JobRequirement], columns: list[Column], *, sheet_name: str, summary: dict | None, metadata: dict | None) -> tuple[bytes, int]:
    """Returns (xlsx bytes, number of rows). Duplicated canonical IDs are skipped defensively."""
    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_name(sheet_name)
    ws.append([c.header for c in columns])
    for i, c in enumerate(columns, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = c.width
    ws.row_dimensions[1].height = 24

    seen: set = set()
    n = 0
    client_counts: dict[str, int] = {}
    for r in rows:
        if r.id in seen:
            continue
        seen.add(r.id)
        n += 1
        ws.append([safe_cell(c.getter(r)) for c in columns])
        row_idx = ws.max_row
        for i, c in enumerate(columns, 1):
            cell = ws.cell(row=row_idx, column=i)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=c.kind in ("wrap", "text"))
            if c.kind == "date":
                cell.number_format = "yyyy-mm-dd"
            elif c.kind == "number":
                cell.number_format = "0"
        client_counts[r.client_name or "(Unknown client)"] = client_counts.get(r.client_name or "(Unknown client)", 0) + 1
    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"

    if summary is not None:
        s = wb.create_sheet("Summary")
        s.append(["Recruitment Requirements Summary"])
        s["A1"].font = Font(bold=True, size=14)
        s.append([])
        s.append(["Total unique requirements", n])
        for k, v in summary.items():
            s.append([safe_cell(k), safe_cell(v)])
        s.append([])
        s.append(["Requirements by client", "Count"])
        for cell in s[s.max_row]:
            cell.font = Font(bold=True)
        for client, cnt in sorted(client_counts.items(), key=lambda kv: -kv[1]):
            s.append([safe_cell(client), cnt])
        s.column_dimensions["A"].width = 40
        s.column_dimensions["B"].width = 24
    if metadata is not None:
        m = wb.create_sheet("Export Info")
        m.append(["Property", "Value"])
        for cell in m[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        for k, v in metadata.items():
            m.append([safe_cell(str(k)), safe_cell(v if not isinstance(v, (list, dict)) else str(v))])
        m.column_dimensions["A"].width = 28
        m.column_dimensions["B"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), n
