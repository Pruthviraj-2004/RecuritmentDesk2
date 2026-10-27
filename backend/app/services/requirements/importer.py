"""Requirement file import: standard XLSX/CSV template, tolerant header matching, row validation, duplicate handling, dry-run preview.

Reuses the hardened file reader of the candidate importer (size/row/column limits, type sniffed from bytes, macro workbooks refused,
formula cells read as values only). Imported requirements go through the same create/update service as manual entry, so they get
versions, normalized fields, skills and fingerprints exactly like any other requirement.
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.exceptions import ValidationFailure
from app.db.types import utcnow
from app.models.enums import LifecycleStatus, ReviewStatus
from app.models.requirements import JobRequirement
from app.schemas.api import RequirementBase
from app.services.candidates import importer as cand_importer
from app.services.normalization.text import requirement_fingerprint
from app.services.requirements import service as svc

MAX_BYTES = cand_importer.MAX_BYTES
MAX_ERROR_REPORT = 200


@dataclass(frozen=True)
class Spec:
    field: str  # RequirementBase field, or "experience_range" (split into minimum/maximum)
    header: str  # standard template header
    kind: str  # text | float | int | list | email | mode | period | range
    example: str
    help: str
    aliases: tuple[str, ...] = ()


SPECS: list[Spec] = [
    Spec("job_title", "job_title", "text", "Senior Python Developer", "Required", ("title", "position", "role", "designation", "job_role", "requirement", "position_title", "opening")),
    Spec("client_name", "client_name", "text", "Globex Corporation", "Recommended. With the title it is used to detect duplicates", ("client", "company", "customer", "account", "end_client", "client_company", "organization", "organisation")),
    Spec("job_reference_id", "job_reference_id", "text", "GX-1042", "Your or the client's requisition / job ID", ("job_id", "req_id", "requisition_id", "reference", "ref_id", "jd_id", "requirement_id", "reference_id", "job_code")),
    Spec("location", "location", "text", "Bengaluru, India", "City / location (or fill city, state and country below)", ("job_location", "work_location", "place")),
    Spec("city", "city", "text", "Bengaluru", "Optional", ()),
    Spec("state", "state", "text", "Karnataka", "Optional", ("province",)),
    Spec("country", "country", "text", "India", "Optional", ()),
    Spec("work_mode", "work_mode", "mode", "hybrid", "remote / hybrid / onsite", ("mode", "work_type", "workplace", "work_model", "remote_hybrid_onsite", "work_from")),
    Spec("experience_range", "experience", "range", "5-8", "Years as a range such as 5-8, or just 5. Or use the two columns below", ("exp", "years_of_experience", "experience_years", "experience_range", "yoe", "exp_range", "total_experience")),
    Spec("minimum_experience", "minimum_experience", "float", "", "Years, 0-60 (alternative to the experience range)", ("min_experience", "min_exp", "experience_min", "exp_min", "minimum_exp")),
    Spec("maximum_experience", "maximum_experience", "float", "", "Years, 0-60", ("max_experience", "max_exp", "experience_max", "exp_max", "maximum_exp")),
    Spec("required_skills", "required_skills", "list", "Python; Django; PostgreSQL; AWS", "Separate with comma, semicolon or |", ("skills", "skill_set", "key_skills", "mandatory_skills", "must_have", "primary_skills", "technologies", "tech_stack", "skillset", "required_skill")),
    Spec("preferred_skills", "preferred_skills", "list", "Docker; Redis", "Nice-to-have skills", ("good_to_have", "nice_to_have", "optional_skills", "secondary_skills", "preferred_skill", "desirable_skills")),
    Spec("salary_min", "salary_min", "float", "2000000", "Number (no currency symbol)", ("min_salary", "salary_from", "minimum_salary", "budget_min", "min_ctc", "ctc_min")),
    Spec("salary_max", "salary_max", "float", "3000000", "Number", ("max_salary", "salary_to", "maximum_salary", "budget_max", "max_ctc", "ctc_max", "budget", "salary", "ctc")),
    Spec("salary_currency", "salary_currency", "text", "INR", "3-letter code", ("currency", "curr")),
    Spec("salary_period", "salary_period", "period", "year", "hour / day / month / year", ("period", "pay_period", "salary_type", "per")),
    Spec("employment_type", "employment_type", "text", "Full-time", "Full-time, Contract, C2H ...", ("type", "job_type", "contract_type", "employment", "engagement")),
    Spec("number_of_openings", "number_of_openings", "int", "2", "Whole number", ("openings", "positions", "vacancies", "no_of_positions", "headcount", "no_of_openings", "number_of_positions")),
    Spec("notice_period", "notice_period", "text", "Immediate to 30 days", "Free text", ("notice", "np")),
    Spec("education", "education", "text", "B.Tech / MCA", "Required qualification", ("qualification", "degree")),
    Spec("certifications", "certifications", "list", "AWS SAA", "Separate with comma, semicolon or |", ("certification", "certs")),
    Spec("department", "department", "text", "Engineering", "Optional", ("dept", "function")),
    Spec("seniority", "seniority", "text", "Senior", "Optional", ("level", "grade")),
    Spec("recruiter_name", "recruiter_name", "text", "Jane Doe", "Optional", ("recruiter", "contact_person", "spoc", "contact_name")),
    Spec("recruiter_email", "recruiter_email", "email", "jane@staffing.com", "Optional; must be a valid email", ("recruiter_mail", "contact_email", "spoc_email", "email", "recruiter_email_id")),
    Spec("recruiter_phone", "recruiter_phone", "text", "+91 98765 43210", "Optional", ("recruiter_mobile", "contact_phone", "phone", "mobile", "contact_number")),
    Spec("submission_email", "submission_email", "email", "", "Where to send profiles, if different", ("send_profiles_to", "submit_to")),
    Spec("relocation", "relocation", "text", "", "Optional", ("relocation_allowed",)),
    Spec("travel", "travel", "text", "", "Optional", ()),
    Spec("job_description", "job_description", "text", "Build and scale our payments APIs.", "Optional, free text", ("description", "jd", "details", "job_details", "summary", "responsibilities")),
]
STANDARD_HEADERS = [s.header for s in SPECS]
_HEADER_MAP: dict[str, Spec] = {}
for _s in SPECS:
    for _k in {_s.header, _s.field, *_s.aliases}:
        _HEADER_MAP.setdefault(_k, _s)
_PERIOD = {"hour": "hour", "hourly": "hour", "hr": "hour", "day": "day", "daily": "day", "month": "month", "monthly": "month", "pm": "month", "per_month": "month",
           "year": "year", "yearly": "year", "annual": "year", "annum": "year", "per_annum": "year", "pa": "year", "yr": "year", "lpa": "year", "annually": "year"}
_MODE = {"remote": "remote", "wfh": "remote", "work_from_home": "remote", "hybrid": "hybrid", "onsite": "onsite", "on_site": "onsite", "office": "onsite",
         "wfo": "onsite", "work_from_office": "onsite", "in_office": "onsite", "unknown": "unknown"}
_EXAMPLE2 = {"job_title": "Data Engineer", "client_name": "Initech", "job_reference_id": "IN-77", "location": "Remote, India", "country": "India", "work_mode": "remote",
             "experience": "4-7", "required_skills": "Python; Spark; SQL; Airflow", "preferred_skills": "Kafka", "salary_min": "1500000", "salary_max": "2200000",
             "salary_currency": "INR", "salary_period": "year", "employment_type": "Contract", "number_of_openings": "1", "notice_period": "30 days", "department": "Data",
             "recruiter_name": "Bob", "recruiter_email": "bob@initech-hr.com", "job_description": "Own the data platform."}


def build_template(fmt: str, examples_rows: bool = True) -> tuple[bytes, str, str]:
    examples = [[s.example for s in SPECS], [_EXAMPLE2.get(s.header, "") for s in SPECS]] if examples_rows else []
    suffix = "" if examples_rows else "-blank"
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(STANDARD_HEADERS)
        w.writerows(examples)
        return ("﻿" + buf.getvalue()).encode("utf-8"), f"requirement-import-template{suffix}.csv", "text/csv"
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    ws.append(STANDARD_HEADERS)
    for row in examples:
        ws.append(row)
    fill = PatternFill("solid", start_color="C2410C")
    for i, h in enumerate(STANDARD_HEADERS, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(16, len(h) + 4)
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Instructions")
    info.append(["Column", "Required", "Format / allowed values", "Example"])
    for c in info[1]:
        c.font = Font(bold=True)
    for s in SPECS:
        info.append([s.header, "Yes" if s.field == "job_title" else "No (optional)", s.help, s.example])
    info.append([])
    info.append(["Only job_title is required. Add client_name too so duplicates can be detected (same client + job ID, or client + title + location)."])
    info.append(["Replace the two example rows with your jobs and keep the header row. Column order does not matter, and extra columns are ignored."])
    info.append(["Duplicates (already in the app, or repeated in the file) are skipped, or you can choose to update the existing requirement with the non-empty cells."])
    for col, w in zip("ABCD", (26, 16, 62, 40)):
        info.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), f"requirement-import-template{suffix}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def map_headers(raw: list[str]) -> tuple[dict[int, Spec], list[str], list[str]]:
    mapping: dict[int, Spec] = {}
    seen: set[str] = set()
    unknown: list[str] = []
    for i, h in enumerate(raw):
        spec = _HEADER_MAP.get(cand_importer.norm_header(h))
        if spec is None:
            if str(h).strip():
                unknown.append(str(h)[:60])
            continue
        if spec.field in seen:
            continue  # first matching column wins
        seen.add(spec.field)
        mapping[i] = spec
    problems = []
    if "job_title" not in {s.field for s in mapping.values()}:
        problems.append("No job title column found. Add a job_title column.")
    return mapping, unknown, problems


@dataclass
class RowResult:
    row: int
    title: str
    client: str | None
    action: str = "create"  # create | update | skip_duplicate | skip_in_file | error
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    values: dict = field(default_factory=dict)
    fingerprint: str | None = None
    matched_id: uuid.UUID | None = None
    matched_title: str | None = None

    def public(self) -> dict:
        return {"row": self.row, "title": self.title, "client": self.client, "action": self.action, "errors": self.errors, "warnings": self.warnings,
                "matched": self.matched_title}


def _key(v: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", v.strip().lower()).strip("_")


def _parse_range(v: str) -> tuple[float | None, float | None]:
    nums = [float(n.replace(",", "")) for n in re.findall(r"\d[\d,]*(?:\.\d+)?", v)]
    if not nums:
        return None, None
    if len(nums) == 1:
        return nums[0], None  # "5" or "5+" means at least 5 years
    return min(nums[0], nums[1]), max(nums[0], nums[1])


def parse_row(row_no: int, cells: list[str], mapping: dict[int, Spec]) -> RowResult:
    raw: dict[str, str] = {}
    for idx, spec in mapping.items():
        v = cells[idx].strip() if idx < len(cells) else ""
        if v:
            raw[spec.field] = v
    r = RowResult(row=row_no, title=(raw.get("job_title") or "")[:120], client=(raw.get("client_name") or None))
    if not raw:
        r.action, r.errors = "error", ["Empty row"]
        return r
    values: dict[str, Any] = {}
    kinds = {s.field: s.kind for s in mapping.values()}
    for f_, v in raw.items():
        kind = kinds[f_]
        if kind == "list":
            values[f_] = cand_importer._split(v)
        elif kind in ("float", "int"):
            n = cand_importer._number(v)
            if n is None:
                r.errors.append(f"{f_}: expected a number, got {v[:30]!r}")
            else:
                values[f_] = int(n) if kind == "int" else n
        elif kind == "mode":
            m = _MODE.get(_key(v))
            if m is None:
                r.warnings.append(f"work_mode {v[:20]!r} not recognised (use remote, hybrid or onsite); left unknown")
            values[f_] = m or "unknown"
        elif kind == "period":
            p = _PERIOD.get(_key(v))
            if p is None:
                r.warnings.append(f"salary_period {v[:20]!r} not recognised (use hour, day, month or year); ignored")
            else:
                values[f_] = p
        elif kind == "range":
            lo, hi = _parse_range(v)
            if lo is None:
                r.errors.append(f"experience: expected years like 5-8, got {v[:30]!r}")
            else:
                values.setdefault("minimum_experience", lo)
                if hi is not None:
                    values.setdefault("maximum_experience", hi)
        else:
            values[f_] = v
    values.pop("experience_range", None)
    if values.get("salary_currency"):
        values["salary_currency"] = str(values["salary_currency"]).strip().upper()[:8]
    if not values.get("job_title"):
        r.errors.append("job_title is required")
    lo, hi = values.get("minimum_experience"), values.get("maximum_experience")
    if lo is not None and hi is not None and lo > hi:
        r.errors.append(f"minimum experience ({lo:g}) is greater than maximum ({hi:g})")
    smin, smax = values.get("salary_min"), values.get("salary_max")
    if smin is not None and smax is not None and smin > smax:
        r.errors.append("salary_min is greater than salary_max")
    if not r.errors:
        try:
            clean = RequirementBase.model_validate(values).model_dump(exclude_none=True)
        except ValidationError as e:
            for err in e.errors()[:5]:
                r.errors.append(f"{'.'.join(str(p) for p in err['loc'])}: {err['msg'][:80]}")
            clean = {}
        values = {**values, **{k: v for k, v in clean.items() if k in values}}
    if r.errors:
        r.action = "error"
        return r
    values["location"] = svc.location_display(values.get("city"), values.get("state"), values.get("country"), values.get("location"))
    r.values = {k: v for k, v in values.items() if v not in (None, "", [])}
    r.fingerprint = requirement_fingerprint(values.get("client_name"), values.get("job_title"), values.get("job_reference_id"), values.get("location"))
    if not values.get("client_name"):
        r.warnings.append("No client_name: duplicates of this job cannot be detected")
    return r


@dataclass
class Analysis:
    total: int
    rows: list[RowResult]
    unknown_headers: list[str]
    recognised: list[str]

    def counts(self) -> dict:
        c = {"create": 0, "update": 0, "skip_duplicate": 0, "skip_in_file": 0, "error": 0}
        for r in self.rows:
            c[r.action] += 1
        return c


def read(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    try:
        return cand_importer.read_rows(filename, data)
    except ValidationFailure as e:
        raise ValidationFailure(str(e.message).replace("candidates", "requirements")) from e


def analyse(db: Session, org_id: uuid.UUID, filename: str, data: bytes, on_duplicate: str) -> Analysis:
    if on_duplicate not in ("skip", "update"):
        raise ValidationFailure("on_duplicate must be 'skip' or 'update'")
    raw_headers, data_rows = read(filename, data)
    mapping, unknown, problems = map_headers(raw_headers)
    if problems:
        raise ValidationFailure(" ".join(problems) + " Download the template to see the expected columns.", code="invalid_template")
    results = [parse_row(i + 2, cells, mapping) for i, cells in enumerate(data_rows)]
    fps = {r.fingerprint for r in results if r.fingerprint}
    existing: dict[str, JobRequirement] = {}
    if fps:
        for req in db.scalars(select(JobRequirement).where(JobRequirement.organization_id == org_id, JobRequirement.canonical_fingerprint.in_(fps),
                                                           JobRequirement.lifecycle_status != LifecycleStatus.MERGED).order_by(JobRequirement.created_at)):
            existing.setdefault(req.canonical_fingerprint, req)
    seen: dict[str, int] = {}
    for r in results:
        if r.action == "error" or not r.fingerprint:
            continue
        if r.fingerprint in seen:
            r.action = "skip_in_file"
            r.warnings.append(f"Same job as row {seen[r.fingerprint]} in this file; skipped")
            continue
        seen[r.fingerprint] = r.row
        match = existing.get(r.fingerprint)
        if match:
            r.matched_id, r.matched_title = match.id, f"{match.job_title or '(untitled)'}" + (f" at {match.client_name}" if match.client_name else "")
            r.action = "update" if on_duplicate == "update" else "skip_duplicate"
    return Analysis(total=len(results), rows=results, unknown_headers=unknown, recognised=sorted({s.header for s in mapping.values()}))


def commit(db: Session, org_id: uuid.UUID, user_id: uuid.UUID, analysis: Analysis, *, approve: bool) -> dict:
    """Create / update requirements. One bad row never aborts the batch (each row runs in a savepoint)."""
    created = updated = skipped = failed = 0
    report: list[dict] = []
    now = utcnow()
    status = ReviewStatus.APPROVED if approve else ReviewStatus.NEEDS_REVIEW
    for r in analysis.rows:
        if r.action in ("error", "skip_in_file", "skip_duplicate"):
            if r.action == "error":
                failed += 1
            else:
                skipped += 1
            if len(report) < MAX_ERROR_REPORT:
                report.append({"row": r.row, "title": r.title, "action": r.action, "problems": r.errors + r.warnings})
            continue
        try:
            with db.begin_nested():
                if r.action == "update" and r.matched_id:
                    req = svc.get_requirement(db, org_id, r.matched_id)
                    patch = {k: v for k, v in r.values.items() if k in svc.EDITABLE_FIELDS}
                    svc.update_requirement(db, req, patch, user_id=user_id, source="import")
                    updated += 1
                else:
                    svc.create_requirement(db, org_id, r.values, user_id=user_id, source="import", review_status=status, received_at=now)
                    created += 1
        except Exception as e:  # noqa: BLE001
            failed += 1
            if len(report) < MAX_ERROR_REPORT:
                report.append({"row": r.row, "title": r.title, "action": "error", "problems": [f"could not be saved ({type(e).__name__})"]})
    db.flush()
    return {"total_rows": analysis.total, "created": created, "updated": updated, "skipped": skipped, "failed": failed, "problems": report}
