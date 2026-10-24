from __future__ import annotations

import io
from datetime import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.candidates import Candidate
from app.services.exports.excel import safe_cell

COLUMNS: list[tuple[str, int, object]] = [
    ("Name", 24, lambda c: c.full_name), ("Email", 28, lambda c: c.email), ("Phone", 16, lambda c: c.phone),
    ("Current title", 26, lambda c: c.current_title), ("Current company", 24, lambda c: c.current_company),
    ("Total exp (yrs)", 12, lambda c: c.total_experience), ("Relevant exp (yrs)", 14, lambda c: c.relevant_experience),
    ("Skills", 42, lambda c: ", ".join(c.skill_names) or None), ("Current location", 20, lambda c: c.current_location),
    ("Preferred locations", 24, lambda c: ", ".join(c.preferred_locations or []) or None),
    ("Willing to relocate", 12, lambda c: None if c.willing_to_relocate is None else ("Yes" if c.willing_to_relocate else "No")),
    ("Notice period", 14, lambda c: c.notice_period), ("Current CTC", 14, lambda c: float(c.current_ctc) if c.current_ctc is not None else None),
    ("Expected CTC", 14, lambda c: float(c.expected_ctc) if c.expected_ctc is not None else None), ("Currency", 9, lambda c: c.currency),
    ("Education", 26, lambda c: c.education), ("Certifications", 26, lambda c: ", ".join(c.certifications or []) or None),
    ("Work authorization", 20, lambda c: c.work_authorization), ("LinkedIn", 30, lambda c: c.linkedin_url), ("Resume link", 30, lambda c: c.resume_url),
    ("Status", 14, lambda c: c.status.replace("_", " ")), ("Rating", 8, lambda c: c.rating), ("Tags", 20, lambda c: ", ".join(c.tags or []) or None),
    ("Source", 14, lambda c: c.source), ("Added", 14, lambda c: c.created_at.astimezone(timezone.utc).replace(tzinfo=None)),
]


def build_workbook(candidates: list[Candidate]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    ws.append([h for h, _, _ in COLUMNS])
    fill = PatternFill("solid", start_color="C2410C")
    for i, (h, w, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for cand in candidates:
        ws.append([safe_cell(fn(cand)) for _, _, fn in COLUMNS])  # untrusted text can never be evaluated as a formula
        ws.cell(row=ws.max_row, column=len(COLUMNS)).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    if candidates:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
