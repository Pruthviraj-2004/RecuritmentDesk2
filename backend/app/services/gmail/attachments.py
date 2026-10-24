"""Safe attachment text extraction. Content is never executed; type is verified from bytes, not names."""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from app.models.enums import AttachmentStatus as S
from app.services.normalization.text import normalize_whitespace

MAX_TEXT_CHARS = 60000
MAX_PDF_PAGES = 40
MAX_ZIP_ENTRIES = 2000
MAX_ZIP_UNCOMPRESSED = 60 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200


@dataclass
class ParsedAttachment:
    status: S
    text: str | None = None
    detected_type: str | None = None
    error: str | None = None


def sniff_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return "zip"
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # legacy .doc/.xls (may contain macros) - not supported
    if data.startswith((b"MZ", b"\x7fELF", b"#!")):
        return "executable"
    if b"\x00" not in data[:4096]:
        try:
            data[:4096].decode("utf-8")
            return "text"
        except UnicodeDecodeError:
            try:
                data[:4096].decode("utf-16")
                return "text"
            except UnicodeDecodeError:
                pass
    return "binary"


def _safe_zip(zf: zipfile.ZipFile) -> str | None:
    infos = zf.infolist()
    if len(infos) > MAX_ZIP_ENTRIES:
        return "archive has too many entries"
    total = 0
    for i in infos:
        total += i.file_size
        if i.compress_size and i.file_size / max(i.compress_size, 1) > MAX_COMPRESSION_RATIO:
            return "suspicious compression ratio"
        if i.filename.startswith(("/", "..")) or ".." in i.filename.split("/"):
            return "unsafe archive path"
    if total > MAX_ZIP_UNCOMPRESSED:
        return "archive expands beyond limits"
    return None


def _parse_docx(data: bytes) -> ParsedAttachment:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return ParsedAttachment(S.FAILED, error="corrupt docx")
    with zf:
        problem = _safe_zip(zf)
        if problem:
            return ParsedAttachment(S.QUARANTINED, error=problem)
        names = set(zf.namelist())
        if "word/document.xml" not in names:
            return ParsedAttachment(S.UNSUPPORTED, error="zip file is not a Word document", detected_type="zip")
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            return ParsedAttachment(S.QUARANTINED, error="document contains macros", detected_type="docx")
    import docx  # python-docx only reads XML; nothing is executed

    try:
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = normalize_whitespace("\n".join(parts))[:MAX_TEXT_CHARS]
    except Exception as e:  # noqa: BLE001
        return ParsedAttachment(S.FAILED, error=f"docx parse error: {type(e).__name__}", detected_type="docx")
    return ParsedAttachment(S.EXTRACTED, text=text, detected_type="docx")


def _parse_pdf(data: bytes) -> ParsedAttachment:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            return ParsedAttachment(S.UNSUPPORTED, error="encrypted pdf", detected_type="pdf")
        out = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            out.append(page.extract_text() or "")
            if sum(map(len, out)) > MAX_TEXT_CHARS:
                break
        text = normalize_whitespace("\n".join(out))[:MAX_TEXT_CHARS]
    except Exception as e:  # noqa: BLE001
        return ParsedAttachment(S.FAILED, error=f"pdf parse error: {type(e).__name__}", detected_type="pdf")
    if not text:
        return ParsedAttachment(S.FAILED, error="no extractable text (scanned PDF? OCR is not enabled)", detected_type="pdf")
    return ParsedAttachment(S.EXTRACTED, text=text, detected_type="pdf")


def parse_attachment(filename: str, declared_type: str | None, data: bytes, max_bytes: int) -> ParsedAttachment:
    if len(data) > max_bytes:
        return ParsedAttachment(S.TOO_LARGE, error=f"larger than {max_bytes // (1024 * 1024)} MB limit")
    if not data:
        return ParsedAttachment(S.FAILED, error="empty file")
    kind = sniff_type(data)
    if kind == "executable":
        return ParsedAttachment(S.QUARANTINED, error="executable content", detected_type=kind)
    if kind == "pdf":
        return _parse_pdf(data)
    if kind == "zip":
        return _parse_docx(data)
    if kind == "text":
        name = (filename or "").lower()
        if name.endswith((".txt", ".text", ".md", ".csv", ".eml")) or (declared_type or "").startswith("text/plain"):
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("utf-16", "replace")
            return ParsedAttachment(S.EXTRACTED, text=normalize_whitespace(text)[:MAX_TEXT_CHARS], detected_type="text")
        return ParsedAttachment(S.UNSUPPORTED, error="unsupported text-like file", detected_type="text")
    return ParsedAttachment(S.UNSUPPORTED, error=f"unsupported file type ({kind})", detected_type=kind)
